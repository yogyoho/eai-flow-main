"""bid_review 解析层测试(EAI-CUSTOM 2026-10-30,spec §2/§4.1)。

覆盖:raw 落盘清单、docx 流序解析+图片 OCR 原位插回、pdf 真页码+扫描页 OCR、
text 切段、OCR 失败隔离(空文本占位不炸任务)、docs_from_entries 全文拼回、
excerpt_window 窗口边界。task_dir monkeypatch 到 tmp_path,不碰真实存储。
"""

import base64
import io

import pytest

from app.extensions.bid_review import parse

# 1x1 PNG(最小合法图片,python-docx add_picture 可收)
_PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")


@pytest.fixture()
def base(tmp_path, monkeypatch):
    monkeypatch.setattr(parse, "task_dir", lambda owner, task: tmp_path / owner / task)
    return tmp_path / "u1" / "t1"


@pytest.fixture()
def ocr_ok():
    async def _fn(name: str, data: bytes, mime: str) -> str:
        return f"OCR[{name}]"

    return _fn


def _make_docx() -> bytes:
    import docx as docx_lib

    d = docx_lib.Document()
    d.add_paragraph("第一段正文")
    d.add_picture(io.BytesIO(_PNG))
    d.add_paragraph("第二段正文")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _make_pdf() -> bytes:
    import fitz

    doc = fitz.open()
    p1 = doc.new_page()
    p1.insert_text((72, 72), "hello pdf page one")
    doc.new_page()  # 空文本层页 → 扫描页
    return doc.tobytes()


@pytest.mark.asyncio
async def test_text_roundtrip_and_windows(base, ocr_ok):
    entry = parse.save_raw("u1", "t1", "tender", "t.txt", b"line1\nline2")
    assert entry["status"] == "raw" and entry["kind"] == "text" and entry["slug"]
    assert (base / "raw" / f"{entry['slug']}.bin").read_bytes() == b"line1\nline2"

    out = await parse.parse_raw("u1", "t1", entry, ocr_ok)
    assert out["status"] == "parsed" and out["chars"] == len("line1\nline2") and out["pages"] == 0

    segs = parse.load_segments("u1", "t1", entry["slug"])
    assert [s["kind"] for s in segs] == ["text"] and segs[0]["page"] is None
    assert parse.docs_from_entries("u1", "t1", [out]) == [{"doc_id": "t.txt", "role": "tender", "text": "line1\nline2"}]

    # excerpt_window:命中段 ±N,首尾夹取
    many = [{"seq": i, "page": None, "kind": "text", "text": str(i), "asset_id": None} for i in range(30)]
    win = parse.excerpt_window(many, 0)
    assert win[0]["seq"] == 0 and len(win) == 11
    assert parse.excerpt_window(many, 15)[0]["seq"] == 5 and parse.excerpt_window(many, 15)[-1]["seq"] == 25


@pytest.mark.asyncio
async def test_docx_flow_order_and_image_ocr(base, ocr_ok):
    entry = parse.save_raw("u1", "t1", "bid", "bid.docx", _make_docx())
    assert entry["kind"] == "docx"
    out = await parse.parse_raw("u1", "t1", entry, ocr_ok)
    assert out["status"] == "parsed" and out["assets"] == 1 and out["pages"] == 0  # D8:docx 无页码

    segs = parse.load_segments("u1", "t1", entry["slug"])
    assert [s["kind"] for s in segs] == ["text", "ocr", "text"]  # 图按原位插回流序
    assert segs[0]["text"] == "第一段正文" and segs[1]["text"].startswith("OCR[asset-")
    assert segs[1]["page"] is None and segs[1]["asset_id"]
    asset_file = next((base / "assets").iterdir())
    assert asset_file.stem == segs[1]["asset_id"]


@pytest.mark.asyncio
async def test_pdf_real_pages_and_scanned_ocr(base, ocr_ok):
    entry = parse.save_raw("u1", "t1", "tender", "t.pdf", _make_pdf())
    out = await parse.parse_raw("u1", "t1", entry, ocr_ok)
    assert out["pages"] == 2 and out["assets"] == 1

    segs = parse.load_segments("u1", "t1", entry["slug"])
    assert segs[0]["kind"] == "text" and segs[0]["page"] == 1 and "hello pdf" in segs[0]["text"]
    assert segs[1]["kind"] == "ocr" and segs[1]["page"] == 2 and segs[1]["asset_id"]  # 扫描页真页码保留


@pytest.mark.asyncio
async def test_ocr_failure_isolated_not_fatal(base):
    """单图 OCR 失败(返回 "")→ 空文本 segment 占位,任务照常 parsed。"""

    async def bad_ocr(name: str, data: bytes, mime: str) -> str:
        return ""

    entry = parse.save_raw("u1", "t1", "bid", "b.docx", _make_docx())
    out = await parse.parse_raw("u1", "t1", entry, bad_ocr)
    assert out["status"] == "parsed"
    segs = parse.load_segments("u1", "t1", entry["slug"])
    assert segs[1]["kind"] == "ocr" and segs[1]["text"] == "" and segs[1]["asset_id"]


@pytest.mark.asyncio
async def test_build_ocr_fn_builtin_fallback(monkeypatch):
    """default_ocr_engine 解析不出(builtin/未配置)→ 内建 eai-flow-ocr callable。"""
    from app.extensions import ocr_resolve

    async def none_entry(db):
        return None

    monkeypatch.setattr(ocr_resolve, "resolve_default_ocr_entry", none_entry)
    fn = await parse.build_ocr_fn(None)
    assert fn is parse._ocr_builtin


@pytest.mark.asyncio
async def test_docs_from_entries_page_markers_and_sharding(base, ocr_ok):
    """抽取 v2(EAI-CUSTOM 2026-10-30):pdf 段带 [[p.N]] 标记(provenance.page 锚点);
    超 shard_chars 按 segment 边界切成同 doc_id 多片段;docx 无页不插标记(D8)。"""
    entry = parse.save_raw("u1", "t1", "tender", "t.pdf", _make_pdf())
    out = await parse.parse_raw("u1", "t1", entry, ocr_ok)
    pieces = parse.docs_from_entries("u1", "t1", [out], shard_chars=40)
    assert len(pieces) >= 2 and all(p["doc_id"] == "t.pdf" for p in pieces)
    assert "[[p.1]]" in pieces[0]["text"]  # pdf 真页码标记

    # docx:无页码 → 无标记,单段小文档不切
    entry2 = parse.save_raw("u1", "t1", "bid", "b.docx", _make_docx())
    out2 = await parse.parse_raw("u1", "t1", entry2, ocr_ok)
    pieces2 = parse.docs_from_entries("u1", "t1", [out2])
    assert len(pieces2) == 1 and "[[p." not in pieces2[0]["text"]
