"""D17 regression contract tests: knowledge upload hook → translation enqueue (T1).

在容器内跑（需真实 postgres；host 连不上 eai-flow-postgres-ext）。
负例=开关关（默认）不入队；正例=开关开建 pending 任务且 lang/tag 正确。
Engine 绑定 session 级循环（pytest-asyncio strict 模式，function 循环会让
全局 asyncpg 引擎跨循环炸）。
"""

import hashlib
import uuid

import pytest
import pytest_asyncio
from sqlalchemy import delete

from app.extensions.database import get_session_factory, init_engine
from app.extensions.models import KnowledgeBase
from app.extensions.translation.enqueue import enqueue_translation_task, should_enqueue
from app.extensions.translation.models import TranslationTask


@pytest_asyncio.fixture(loop_scope="session", scope="session")
async def _engine():
    await init_engine()
    yield


async def _cleanup(db, kb_id) -> None:
    """显式序清理：先任务后 KB（FK 方向），bulk delete 不依赖 ORM 级联排序。"""
    await db.execute(delete(TranslationTask).where(TranslationTask.kb_id == kb_id))
    await db.execute(delete(KnowledgeBase).where(KnowledgeBase.id == kb_id))
    await db.commit()


def _make_docx(path, text="The annual report of the mining company for the fiscal year."):
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    doc.save(str(path))
    return path


async def _unique_kb(db, enabled: bool) -> KnowledgeBase:
    from sqlalchemy import select

    from app.extensions.models import User

    owner_id = (await db.execute(select(User.id).limit(1))).scalar()
    if owner_id is None:
        pytest.skip("no users row in extensions db to own the test KB")
    kb = KnowledgeBase(
        name=f"tlt-hook-test-{uuid.uuid4().hex[:8]}",
        owner_id=owner_id,
        kb_type="ragflow",
        translation_enabled=enabled,
    )
    db.add(kb)
    await db.commit()
    await db.refresh(kb)
    return kb


@pytest.mark.asyncio(loop_scope="session")
async def test_d17_switch_off_produces_no_enqueue(_engine, tmp_path):
    """开关关（schema 默认）→ should_enqueue False → 钩子分支零行为。"""
    async with get_session_factory()() as db:
        kb = await _unique_kb(db, enabled=False)
        kb_id = kb.id
        try:
            assert should_enqueue(kb) is False
        finally:
            await _cleanup(db, kb_id)


@pytest.mark.asyncio(loop_scope="session")
async def test_switch_on_enqueues_pending_task_with_lang_tag(_engine, tmp_path):
    """开关开 → 后台解析（synthetic docx）→ pending 任务，lang=en / tag=narrative。"""
    p = _make_docx(tmp_path / "hook.docx")
    sha = hashlib.sha256(p.read_bytes()).hexdigest()
    async with get_session_factory()() as db:
        kb = await _unique_kb(db, enabled=True)
        kb_id = kb.id
        try:
            task_id = await enqueue_translation_task(
                kb_id=kb_id,
                document_id=None,
                doc_sha=sha,
                source_uri=str(p),
                file_path=str(p),
            )
            assert task_id is not None
            task = await db.get(TranslationTask, uuid.UUID(task_id))
            assert task is not None
            assert task.state == "pending"
            assert task.source_lang == "en"
            assert task.content_tag == "narrative"
            assert task.quality_tier == "draft"
        finally:
            await _cleanup(db, kb_id)


@pytest.mark.asyncio(loop_scope="session")
async def test_enqueue_survives_unparseable_file(_engine, tmp_path):
    """解析失败不阻塞入队（lang=mixed 兜底，worker 阶段再判）——设计构件 4。"""
    bad = tmp_path / "broken.docx"
    bad.write_bytes(b"not a real docx")
    sha = hashlib.sha256(bad.read_bytes()).hexdigest()
    async with get_session_factory()() as db:
        kb = await _unique_kb(db, enabled=True)
        kb_id = kb.id
        try:
            task_id = await enqueue_translation_task(kb_id=kb_id, document_id=None, doc_sha=sha, source_uri=str(bad), file_path=str(bad))
            task = await db.get(TranslationTask, uuid.UUID(task_id))
            assert task.source_lang == "mixed"
        finally:
            await _cleanup(db, kb_id)


@pytest.mark.asyncio(loop_scope="session")
async def test_delete_knowledge_doc_archives_translation_tasks(_engine, monkeypatch):
    """P1-5: 知识库删文档 → 翻译任务归档、中文份下线、document_id 释放。"""
    from app.extensions.knowledge import service as knowledge_service
    from app.extensions.knowledge.service import DocumentService
    from app.extensions.models import Document

    class _FakeClient:
        def __init__(self):
            self.deleted: list[tuple[str, str]] = []

        async def delete_document(self, dataset_id: str, document_id: str) -> None:
            self.deleted.append((dataset_id, document_id))

    fake = _FakeClient()

    async def _fake_service_client(db):
        return fake

    monkeypatch.setattr(knowledge_service, "get_service_client", _fake_service_client)

    async with get_session_factory()() as db:
        kb = await _unique_kb(db, enabled=True)
        kb.ragflow_dataset_id = "ds-hook"
        doc = Document(knowledge_base_id=kb.id, name="a.docx", file_path="/tmp/not-real.docx", ragflow_document_id="doc-1")
        db.add(doc)
        await db.commit()
        await db.refresh(doc)
        task = TranslationTask(
            kb_id=kb.id,
            document_id=doc.id,
            doc_sha=uuid.uuid4().hex + uuid.uuid4().hex,
            source_uri="/tmp/not-real.docx",
            source_lang="en",
            state="delivered",
            quality_tier="draft",
            progress=0,
            retry_count=0,
            task_version=1,
            ragflow_zh_doc_id="zh-1",
        )
        db.add(task)
        await db.commit()
        task_id, doc_id = task.id, doc.id
        try:
            await DocumentService.delete_doc(db, doc)
            await db.refresh(task)
            assert task.state == "archived"
            assert task.document_id is None
            assert ("ds-hook", "doc-1") in fake.deleted
            assert ("ds-hook", "zh-1") in fake.deleted
            assert await db.get(Document, doc_id) is None
        finally:
            await db.execute(delete(TranslationTask).where(TranslationTask.id == task_id))
            await _cleanup(db, kb.id)


@pytest.mark.asyncio(loop_scope="session")
async def test_zh_document_is_not_enqueued(_engine, tmp_path):
    """P2-14: knowledge 钩子遇到中文文档不建翻译任务。"""
    p = _make_docx(tmp_path / "zh.docx", text="这是一份中文技术资料，不需要翻译。")
    sha = hashlib.sha256(p.read_bytes()).hexdigest()
    async with get_session_factory()() as db:
        kb = await _unique_kb(db, enabled=True)
        kb_id = kb.id
        try:
            task_id = await enqueue_translation_task(
                kb_id=kb_id,
                document_id=None,
                doc_sha=sha,
                source_uri=str(p),
                file_path=str(p),
            )
            assert task_id is None
        finally:
            await _cleanup(db, kb_id)
