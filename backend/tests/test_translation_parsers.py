"""Parser unit tests — synthetic docs built in-test, no external fixtures."""

import pytest

from app.extensions.translation.langdetect import detect_language
from app.extensions.translation.parsers import (
    Block,
    OcrRequired,
    blocks_text,
    classify_tag,
    parse_document,
)


def _make_docx(path, paragraphs, table_rows=None):
    from docx import Document

    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_rows:
        t = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, val in enumerate(row):
                t.rows[r].cells[c].text = val
    doc.save(str(path))
    return path


def _make_xlsx(path, sheet_rows, sheet_name="Sheet1"):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in sheet_rows:
        ws.append(row)
    wb.save(str(path))
    return path


def _make_pptx(path, slide_texts):
    from pptx import Presentation

    prs = Presentation()
    for text in slide_texts:
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # title-only layout
        body = slide.shapes.placeholders[0]
        body.text_frame.text = text
    prs.save(str(path))
    return path


class TestDocxParser:
    def test_paragraphs_and_table_in_order(self, tmp_path):
        p = _make_docx(
            tmp_path / "a.docx",
            paragraphs=["Annual production report", "Iron ore output increased."],
            table_rows=[["Year", "Tonnes"], ["2024", "3,200"]],
        )
        blocks = parse_document(p)
        types = [b.type for b in blocks]
        assert "paragraph" in types and "table" in types
        assert types.index("paragraph") < types.index("table")
        table = next(b for b in blocks if b.type == "table")
        assert table.meta["cells"][0] == ["Year", "Tonnes"]

    def test_english_docx_detected_as_en(self, tmp_path):
        p = _make_docx(tmp_path / "en.docx", paragraphs=["The report of the company for the year."])
        assert detect_language(blocks_text(parse_document(p))) == "en"


class TestXlsxParser:
    def test_sheet_header_and_table(self, tmp_path):
        p = _make_xlsx(tmp_path / "a.xlsx", sheet_rows=[["Item", "Qty"], ["Ore", "1200"]])
        blocks = parse_document(p)
        assert blocks[0].type == "heading" and blocks[0].text == "Sheet1"
        table = blocks[1]
        assert table.type == "table"
        assert table.meta["cells"][1] == ["Ore", "1200"]


class TestPptxParser:
    def test_slide_text(self, tmp_path):
        p = _make_pptx(tmp_path / "a.pptx", slide_texts=["Board meeting Q3", "Revenue is stable"])
        blocks = parse_document(p)
        assert len(blocks) == 2
        assert {b.meta["slide"] for b in blocks} == {1, 2}
        assert {b.text for b in blocks} == {"Board meeting Q3", "Revenue is stable"}


class TestPdfRouting:
    def test_born_digital_pdf_parses(self, tmp_path):
        import fitz

        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "The annual report of the mining company for the fiscal year.")
        doc.save(str(tmp_path / "a.pdf"))
        doc.close()
        blocks = parse_document(tmp_path / "a.pdf")
        assert blocks and blocks[0].meta["page"] == 1

    def test_image_raises_ocr_required(self, tmp_path):
        (tmp_path / "scan.png").write_bytes(b"\x89PNG fake")
        with pytest.raises(OcrRequired):
            parse_document(tmp_path / "scan.png")

    def test_unsupported_ext_raises_value(self, tmp_path):
        (tmp_path / "a.xyz").write_text("x")
        with pytest.raises(ValueError):
            parse_document(tmp_path / "a.xyz")


class TestClassifyTag:
    def test_no_tables_is_narrative(self):
        assert classify_tag([Block("b1", "paragraph", "text")]) == "narrative"

    def test_table_heavy_is_tabular(self):
        blocks = [Block(f"b{i}", "table", "x\ty", meta={"cells": [["x", "y"]]}) for i in range(4)]
        blocks.append(Block("b9", "paragraph", "note"))
        assert classify_tag(blocks) == "tabular"

    def test_some_tables_is_mixed(self):
        # 1/5 = 20% ≤ 30% 阈值 → mixed（既有表但非表格主导）
        blocks = [Block(f"b{i}", "paragraph", "a") for i in range(4)]
        blocks.append(Block("b9", "table", "c\td", meta={"cells": [["c", "d"]]}))
        assert classify_tag(blocks) == "mixed"

    def test_empty_is_none(self):
        assert classify_tag([]) is None
